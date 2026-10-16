"""Stage 11: Excel human-review export.

Per DATA_PIPELINE_MVP_SPEC.md §4.7: every job in the run stays locatable, with
identity, provenance, eligibility/failure state, HD/Fit evidence, unknowns, and
empty human-editable columns. Column contract is only P1 (Missing P1 in
MVP_WORKFLOW_STAGE_CONTRACT_MAP.md row 11) -- never treated as the system of
record, no re-import behavior implemented.

Tabs follow the exp-003 linkedin workbook layout so a reviewer reads only what
each group of jobs needs:
  Summary        run counts and Fit label distribution
  Fit_Passed     jobs that passed HD, ranked by Fit label, with derived reasons
  HD_Excluded    jobs HD excluded, with a plain-language failed_reason
  Fit Skill Map  one row per requirement (CAPABILITY_FIT_POLICY.md §9.1)
  Not_Evaluated  blocked inputs and failed evaluations
  Run_Metadata   the run manifest
In both job tabs the eight per-rule HD columns are a collapsed column group.

Re-export an existing run from the JSONL stores (no model calls):
    python -m pipeline.review_export <run_id> [--force]
"""

from __future__ import annotations

import argparse
import json
from collections import Counter
from pathlib import Path

from openpyxl import Workbook
from openpyxl.cell.rich_text import CellRichText, TextBlock
from openpyxl.cell.text import InlineFont
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.worksheet import Worksheet

from pipeline.assessment_schema import FitEvidence, EvaluationResult
from pipeline.job_schema import Job

# Keyed by rule_id (HARD_DISQUALIFIER_POLICY.md §3): the model-written
# rule_name is not stable across results ("compensation" vs "compensation_threshold").
# rule_id -> (column header, short label for hd_to_confirm, failed_reason prefix)
HD_RULES = {
    "3.1(a)": ("hd_entry_program", "Entry-level programme", "Entry-level programme (internship/trainee/graduate)"),
    "3.1(b)": ("hd_senior_years", "Years of experience", "Requires too many years of experience"),
    "3.1(c)": ("hd_people_mgmt", "People management", "Requires managing a team"),
    "3.2": ("hd_compensation", "Compensation", "Pay below threshold"),
    "3.3": ("hd_contract", "Contract type", "Unsuitable contract type"),
    "3.4": ("hd_language", "Language", "Language requirement not met"),
    "3.5": ("hd_credentials", "Credentials", "Requires a specific credential"),
    "3.6": ("hd_field_travel", "Field work/travel", "Mainly field work or travel"),
}
HD_RULE_COLUMNS = [v[0] for v in HD_RULES.values()]

FIT_PASSED_COLUMNS = [
    "job_id", "company", "title", "source_url",
    *HD_RULE_COLUMNS, "hd_to_confirm",
    "fit_label", "core_matches", "key_gap", "to_confirm", "industry",
    "your_fit_label", "review_notes",
    "evaluated_at", "model",
]
HD_EXCLUDED_COLUMNS = [
    "job_id", "company", "title", "source_url",
    *HD_RULE_COLUMNS, "failed_reason",
    "review_decision", "review_notes",
    "evaluated_at", "model",
]
NOT_EVALUATED_COLUMNS = [
    "job_id", "company", "title", "source_url",
    "evaluation_status", "failure_stage", "reason",
]
_WIDTHS = {
    "company": 22, "title": 34, "source_url": 14, "hd_to_confirm": 26,
    "fit_label": 13, "core_matches": 55, "key_gap": 55, "to_confirm": 50,
    "industry": 40, "failed_reason": 60, "your_fit_label": 14,
    "review_decision": 16, "review_notes": 30, "reason": 60,
}
_WRAP_COLUMNS = {"title", "hd_to_confirm", "core_matches", "key_gap", "to_confirm",
                 "industry", "failed_reason", "review_notes", "reason"}

FIT_LABEL_ORDER = ["StrongFit", "PlausibleFit", "WeakFit", "Unclear"]
REVIEW_DECISIONS = ["Agree to exclude", "Should keep", "Right outcome, wrong reason", "Can't tell"]

_RULE_FILLS = {
    "Pass": PatternFill("solid", fgColor="E2F1E8"),
    "Fail": PatternFill("solid", fgColor="F7D9D6"),
    "Flag": PatternFill("solid", fgColor="F9E7C6"),
    "Unknown": PatternFill("solid", fgColor="E9ECEA"),
}
_LABEL_FILLS = {
    "StrongFit": PatternFill("solid", fgColor="CDEBD8"),
    "PlausibleFit": PatternFill("solid", fgColor="D6E6F6"),
    "WeakFit": PatternFill("solid", fgColor="F9E7C6"),
    "Unclear": PatternFill("solid", fgColor="E6DDF3"),
}
_HEADER_FILL = PatternFill("solid", fgColor="EEF0EC")

SKILL_MAP_COLUMNS = [
    "job_id", "company", "title", "fit_label", "skill_id",
    "requirement", "your_evidence", "assessment", "state",
]
_SKILL_MAP_WIDTHS = {"requirement": 40, "your_evidence": 55, "assessment": 70}

# Policy §9.1: the verdict is fixed per state; the model writes only the reasoning.
STATE_VERDICTS = {
    "done": "Direct evidence",
    "participated": "Partial match (contributed, not owned)",
    "adjacent": "Transferable",
    "unconfirmed": "Not shown in CV (to confirm)",
    "not_done": "Confirmed gap",
}
_BOLD = InlineFont(b=True)


def _verdict(e: FitEvidence) -> str:
    if e.state == "adjacent":
        return f"{STATE_VERDICTS['adjacent']} (shifted: {', '.join(e.shifted_dimensions)})"
    return STATE_VERDICTS[e.state]


def _your_evidence(e: FitEvidence) -> str:
    if e.state == "unconfirmed":
        return "None found in CV"
    text = e.candidate_evidence or ""
    if e.scope:
        text += f"\nScope: {e.scope}"
    return text


def _assessment(e: FitEvidence) -> CellRichText:
    """`**Verdict.** reasoning`, as in the policy §9.1 example table."""
    return CellRichText(TextBlock(_BOLD, f"{_verdict(e)}."), f" {e.reasoning}")


def _skill_map_rows(result: EvaluationResult) -> list[list]:
    """Policy §9.1: core skills most central first, then the rest in Layer 1
    order, then the industry. Jobs without a Fit result get no rows."""
    fit = result.capability_fit
    if fit.label is None:
        return []
    prefix = [result.job_id, result.company, result.title, fit.label]
    core_order = [c.skill_id for c in fit.core_skills]
    skills = sorted(fit.skills, key=lambda s: core_order.index(s.skill_id)
                    if s.skill_id in core_order else len(core_order))
    evidence = {e.skill_id: e for e in fit.evidence}
    rows = []
    for s in skills:
        e = evidence[s.skill_id]
        star = "★ " if s.skill_id in core_order else ""
        rows.append(prefix + [s.skill_id, f"{star}{s.skill} ({s.jd_signal})",
                              _your_evidence(e), _assessment(e), e.state])
    if fit.industry and fit.industry.industry and fit.industry_evidence:
        e = fit.industry_evidence
        rows.append(prefix + ["industry",
                              f"Industry: {fit.industry.industry} ({fit.industry.jd_signal})",
                              _your_evidence(e), _assessment(e), e.state])
    return rows


def _rule_states(result: EvaluationResult) -> list[str | None]:
    by_id = {r.rule_id: r.state for r in result.hard_disqualifier.rules}
    return [by_id.get(rule_id) for rule_id in HD_RULES]


def _failed_reason(result: EvaluationResult) -> str:
    """Why HD excluded the job, in plain words plus the JD quote -- no rule ids."""
    lines = []
    for r in result.hard_disqualifier.rules:
        if r.state != "Fail":
            continue
        prefix = HD_RULES.get(r.rule_id, (None, None, r.rule_name))[2]
        lines.append(f"{prefix}: “{r.evidence_quote}”" if r.evidence_quote else f"{prefix}: {r.reason}")
    return "\n".join(lines)


def _hd_to_confirm(result: EvaluationResult) -> str:
    parts = []
    for r in result.hard_disqualifier.rules:
        label = HD_RULES.get(r.rule_id, (None, r.rule_name))[1]
        if r.state == "Flag":
            parts.append(f"{label} (borderline)")
        elif r.state == "Unknown":
            parts.append(f"{label} (not stated in JD)")
    return "; ".join(parts)


def _core_lines(result: EvaluationResult, states: set[str], *, with_evidence: bool) -> str:
    """One bullet per core skill in `states`, most central first."""
    fit = result.capability_fit
    skills = {s.skill_id: s for s in fit.skills}
    evidence = {e.skill_id: e for e in fit.evidence}
    lines = []
    for c in fit.core_skills:
        e = evidence.get(c.skill_id)
        if e is None or e.state not in states:
            continue
        detail = e.candidate_evidence if with_evidence and e.candidate_evidence else e.reasoning
        lines.append(f"• {skills[c.skill_id].skill} [{_verdict(e)}] — {detail}")
    return "\n".join(lines)


def _to_confirm(result: EvaluationResult) -> str:
    lines = [_core_lines(result, {"unconfirmed"}, with_evidence=False)]
    lines += [f"• {u}" for u in result.capability_fit.unknowns]
    return "\n".join(line for line in lines if line)


def _industry(result: EvaluationResult) -> str:
    fit = result.capability_fit
    if fit.industry is None:
        return ""
    if fit.industry.industry is None:
        return "not revealed by JD"
    text = f"{fit.industry.industry} ({fit.industry.jd_signal})"
    e = fit.industry_evidence
    if e is not None:
        text += f" [{_verdict(e)}] — {e.reasoning}"
    return text


def _model(result: EvaluationResult) -> str:
    return f"{result.model_identifier} / {result.prompt_version}"


def _fit_passed_row(result: EvaluationResult) -> list:
    return [
        result.job_id, result.company, result.title, result.source_url,
        *_rule_states(result), _hd_to_confirm(result),
        result.capability_fit.label,
        _core_lines(result, {"done", "participated"}, with_evidence=True),
        _core_lines(result, {"adjacent", "not_done"}, with_evidence=False),
        _to_confirm(result),
        _industry(result),
        "", "",  # your_fit_label, review_notes -- human fills in
        result.evaluated_at, _model(result),
    ]


def _hd_excluded_row(result: EvaluationResult) -> list:
    return [
        result.job_id, result.company, result.title, result.source_url,
        *_rule_states(result), _failed_reason(result),
        "", "",  # review_decision, review_notes -- human fills in
        result.evaluated_at, _model(result),
    ]


def _not_evaluated_row(result: EvaluationResult) -> list:
    reason = "; ".join(result.blocking_reasons) or (result.failure_reason or "").split("\n")[0]
    return [result.job_id, result.company, result.title, result.source_url,
            result.evaluation_status, result.failure_stage, reason]


def _partition(results: list[EvaluationResult]):
    passed, excluded, not_evaluated = [], [], []
    for r in results:
        if r.hard_disqualifier.overall_decision == "Exclude":
            excluded.append(r)
        elif r.evaluation_status == "Completed" and r.capability_fit.label:
            passed.append(r)
        else:
            not_evaluated.append(r)
    passed.sort(key=lambda r: FIT_LABEL_ORDER.index(r.capability_fit.label))
    return passed, excluded, not_evaluated


def _write_table(ws: Worksheet, columns: list[str], rows: list[list]) -> None:
    ws.append(columns)
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.fill = _HEADER_FILL
    for row in rows:
        ws.append(row)
    ws.freeze_panes = "D2"  # keep job_id/company/title in view
    ws.auto_filter.ref = ws.dimensions
    wrap = Alignment(wrap_text=True, vertical="top")
    top = Alignment(vertical="top")
    for idx, name in enumerate(columns, start=1):
        letter = get_column_letter(idx)
        ws.column_dimensions[letter].width = _WIDTHS.get(name, 12)
        for (cell,) in ws.iter_rows(min_row=2, min_col=idx, max_col=idx):
            cell.alignment = wrap if name in _WRAP_COLUMNS else top
            if name == "source_url" and cell.value:
                cell.hyperlink = cell.value
                cell.value = "JD ↗"
                cell.font = Font(color="2B5A84", underline="single")
            elif name in HD_RULE_COLUMNS:
                cell.fill = _RULE_FILLS.get(cell.value, PatternFill())
            elif name == "fit_label":
                cell.fill = _LABEL_FILLS.get(cell.value, PatternFill())


def _collapse_rule_columns(ws: Worksheet, columns: list[str]) -> None:
    first = columns.index(HD_RULE_COLUMNS[0]) + 1
    last = columns.index(HD_RULE_COLUMNS[-1]) + 1
    ws.column_dimensions.group(get_column_letter(first), get_column_letter(last),
                               hidden=True, outline_level=1)


def _dropdown(ws: Worksheet, columns: list[str], column: str, options: list[str], n_rows: int) -> None:
    """A suggestion list, not a constraint: free text is still accepted."""
    if n_rows == 0:
        return
    letter = get_column_letter(columns.index(column) + 1)
    dv = DataValidation(type="list", formula1='"' + ",".join(options) + '"',
                        allow_blank=True, showErrorMessage=False)
    dv.add(f"{letter}2:{letter}{n_rows + 1}")
    ws.add_data_validation(dv)


def _read_json(path: Path) -> dict:
    return json.loads(path.read_text(encoding="utf-8")) if path.exists() else {}


class HumanReviewExporter:
    def export(self, results: list[EvaluationResult], jobs_by_id: dict[str, Job],
               path: Path) -> Path:
        path.parent.mkdir(parents=True, exist_ok=True)
        passed, excluded, not_evaluated = _partition(results)
        manifest = _read_json(path.parent / "run_manifest.json")
        counts = _read_json(path.parent / "run_summary.json").get("counts", {})

        wb = Workbook()
        summary = wb.active
        summary.title = "Summary"
        labels = Counter(r.capability_fit.label for r in passed)
        summary_rows = [
            ("Run", manifest.get("run_id") or (results[0].run_id if results else "")),
            ("Started at", manifest.get("started_at", "")),
            ("Hard-disqualifier policy", manifest.get("hard_disqualifier_policy_version", "")),
            ("Capability-fit policy", manifest.get("capability_fit_policy_version", "")),
            ("Candidate profile", manifest.get("candidate_profile_version", "")),
            ("Jobs captured", counts.get("captured", "")),
            ("New or changed jobs", counts.get("new", 0) + counts.get("changed", 0) if counts else ""),
            ("Jobs evaluated in this run", len(results)),
            ("HD excluded (tab HD_Excluded)", len(excluded)),
            ("Passed HD with Fit result (tab Fit_Passed)", len(passed)),
            *[(f"  {label}", labels.get(label, 0)) for label in FIT_LABEL_ORDER],
            ("Not evaluated (tab Not_Evaluated)", len(not_evaluated)),
        ]
        summary.append(["Metric", "Value"])
        for cell in summary[1]:
            cell.font = Font(bold=True)
            cell.fill = _HEADER_FILL
        for row in summary_rows:
            summary.append(list(row))
        summary.column_dimensions["A"].width = 42
        summary.column_dimensions["B"].width = 44

        ws = wb.create_sheet("Fit_Passed")
        _write_table(ws, FIT_PASSED_COLUMNS, [_fit_passed_row(r) for r in passed])
        _collapse_rule_columns(ws, FIT_PASSED_COLUMNS)
        _dropdown(ws, FIT_PASSED_COLUMNS, "your_fit_label", FIT_LABEL_ORDER, len(passed))

        ws = wb.create_sheet("HD_Excluded")
        _write_table(ws, HD_EXCLUDED_COLUMNS, [_hd_excluded_row(r) for r in excluded])
        _collapse_rule_columns(ws, HD_EXCLUDED_COLUMNS)
        _dropdown(ws, HD_EXCLUDED_COLUMNS, "review_decision", REVIEW_DECISIONS, len(excluded))

        skill_map = wb.create_sheet("Fit Skill Map")
        skill_map.append(SKILL_MAP_COLUMNS)
        for cell in skill_map[1]:
            cell.font = Font(bold=True)
        skill_map.freeze_panes = "A2"
        for result in passed:
            for row in _skill_map_rows(result):
                skill_map.append(row)
        wrap = Alignment(wrap_text=True, vertical="top")
        for idx, name in enumerate(SKILL_MAP_COLUMNS, start=1):
            if name in _SKILL_MAP_WIDTHS:
                skill_map.column_dimensions[skill_map.cell(1, idx).column_letter].width = _SKILL_MAP_WIDTHS[name]
                for (cell,) in skill_map.iter_rows(min_row=2, min_col=idx, max_col=idx):
                    cell.alignment = wrap

        ws = wb.create_sheet("Not_Evaluated")
        _write_table(ws, NOT_EVALUATED_COLUMNS, [_not_evaluated_row(r) for r in not_evaluated])

        meta = wb.create_sheet("Run_Metadata")
        meta.append(["Field", "Value"])
        for cell in meta[1]:
            cell.font = Font(bold=True)
            cell.fill = _HEADER_FILL
        for key, value in manifest.items():
            meta.append([key, ", ".join(value) if isinstance(value, list) else value])
        meta.column_dimensions["A"].width = 34
        meta.column_dimensions["B"].width = 60

        wb.save(path)
        return path


def main() -> None:
    from pipeline.result_repository import EvaluationRepository, JobRepository

    parser = argparse.ArgumentParser(description="Re-export a run's review workbook from JSONL.")
    parser.add_argument("run_id")
    parser.add_argument("--force", action="store_true", help="overwrite an existing review_export.xlsx")
    args = parser.parse_args()

    path = Path(f"data/runs/{args.run_id}/review_export.xlsx")
    if path.exists() and not args.force:
        raise SystemExit(f"{path} exists (it may hold review notes); pass --force to overwrite")
    results = EvaluationRepository().read_all_for_run(args.run_id)
    if not results:
        raise SystemExit(f"no evaluation results for run {args.run_id}")
    HumanReviewExporter().export(results, JobRepository().read_current(), path)
    print(f"Human review export: {path}")


if __name__ == "__main__":
    main()
