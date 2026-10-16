"""Offline end-to-end smoke test for the whole orchestrator wiring.

Zero network/API calls: a fake SourceAdapter returns two canned SourceRecords
(one that Hard-Disqualifier-Fails, one that proceeds to a Fit label), and the
JobExtractor/HardDisqualifierEvaluator/CapabilityFitEvaluator clients are
stubbed. This is the last check before spending any real money -- it exercises
the exact same code path run_pipeline.py uses, including the run_id
consistency between raw capture, Job/evaluation persistence, and
reconciliation (a real bug caught here: RunContext.new() was minting a
second, different run_id from the manifest's).

Only paths this test creates are removed afterward -- confirmed empty
beforehand (data/raw, data/jobs, data/evaluations, data/runs do not exist yet
in this repo).

Run: python -m pipeline.validate_phase_d
"""

from __future__ import annotations

import json
import shutil
from pathlib import Path

from openpyxl import load_workbook

import pipeline.orchestrator as orchestrator_module
from pipeline.candidate_profile import CandidateProfile
from pipeline.orchestrator import PipelineOrchestrator
from pipeline.query_catalog import QueryCatalog
from pipeline.run_manifest import RunManifest
from pipeline.source_adapter import LocationFilter, QueryDefinition, RunContext, SourceAdapter
from pipeline.source_record import SourceRecord, compute_content_hash, utcnow_iso

problems: list[str] = []


def check(condition: bool, message: str) -> None:
    if not condition:
        problems.append(message)


class FakeTextBlock:
    def __init__(self, text: str):
        self.type = "text"
        self.text = text


class FakeResponse:
    def __init__(self, text: str):
        self.content = [FakeTextBlock(text)]


class FakeMessages:
    def __init__(self, response_fn):
        self._response_fn = response_fn

    def create(self, **kwargs):
        return FakeResponse(self._response_fn(kwargs))


class FakeClient:
    def __init__(self, response_fn):
        self.messages = FakeMessages(response_fn)


def _fake_source_record(source_url: str, source_job_id: str) -> SourceRecord:
    return SourceRecord(
        source_record_id=f"rec-{source_job_id}", source="Indeed", source_job_id=source_job_id,
        source_url=source_url, resolved_url=source_url,
        collection_run_id="placeholder", sourcing_config_version="sourcing-config-v0.1",
        source_config_id="indeed-de-v0.1", query_group_id="customer-success-v0.1",
        query_id="cs-customer-success-manager-en-v0.1",
        query_text_or_parameters="Customer Success Manager", captured_at=utcnow_iso(),
        adapter_version="fake-v0.1", parser_version="fake-v0.1", capture_method="Fake",
        raw_title="Customer Success Manager", raw_company="Acme",
        raw_location={"city": "München", "country": "Deutschland"},
        raw_description="A real JD would be here.", raw_publication_date="2026-09-01",
        raw_job_status=False, raw_work_mode=False, raw_payload_reference="inline",
        content_hash=compute_content_hash(title="Customer Success Manager", company="Acme",
                                           location="Munich", description="A real JD would be here.",
                                           publication_date="2026-09-01"),
        collection_status="Success",
    )


class FakeAdapter(SourceAdapter):
    source = "Indeed"
    source_config_id = "indeed-de-v0.1"
    adapter_version = "fake-v0.1"
    parser_version = "fake-v0.1"
    capture_method = "Fake"

    def __init__(self, client=None):
        self._called = False

    def capture(self, query: QueryDefinition, location: LocationFilter,
                run_context: RunContext, max_rows: int) -> list[SourceRecord]:
        if self._called:
            return []
        self._called = True
        return [
            _fake_source_record("https://example.com/job/test-pass", "test-job-pass"),
            _fake_source_record("https://example.com/job/test-fail", "test-job-fail"),
        ]


def _fake_extraction_response(kwargs) -> str:
    return json.dumps({
        "description": "A role.", "responsibilities": "Manage accounts.",
        "qualifications_requirements": "3 years experience.",
        "description_status": "Full", "work_mode": "Hybrid",
    })


def _fake_hd_response(kwargs) -> str:
    user_content = kwargs["messages"][0]["content"]
    if "test-fail" in user_content:
        return json.dumps({
            "rules": [{"rule_id": "3.2", "rule_name": "compensation", "state": "Fail",
                       "evidence_quote": "EUR 30,000", "candidate_evidence": None,
                       "reason": "Below the EUR 50k threshold."}],
            "overall_decision": "Exclude",
        })
    return json.dumps({
        "rules": [{"rule_id": "3.2", "rule_name": "compensation", "state": "Unknown",
                   "evidence_quote": None, "candidate_evidence": None, "reason": "Not disclosed."}],
        "overall_decision": "ProceedToCapabilityFit",
    })


def _fake_fit_response(kwargs) -> str:
    # Capability Fit v0.3 makes two calls; tell them apart by requested schema.
    schema_keys = set(kwargs["output_config"]["format"]["schema"]["properties"])
    if "skills" in schema_keys:
        return json.dumps({
            "skills": [{"skill_id": "S1", "skill": "B2B account management", "category": "method",
                        "jd_signal": "required", "jd_quote": "Manage accounts."},
                       {"skill_id": "S2", "skill": "Customer onboarding", "category": "method",
                        "jd_signal": "responsibility", "jd_quote": "Onboard customers."}],
            "core_skills": [{"skill_id": "S1", "reason": "Central function."},
                            {"skill_id": "S2", "reason": "Primary responsibility."}],
            "industry": {"industry": None, "jd_signal": "context_only", "jd_quote": None},
        })
    return json.dumps({
        "evidence": [
            {"skill_id": "S1", "state": "done", "candidate_evidence": "ev-01: managed pharma B2B accounts",
             "scope": "Owned two named accounts.", "shifted_dimensions": [],
             "reasoning": "Owned B2B accounts end-to-end."},
            {"skill_id": "S2", "state": "adjacent", "candidate_evidence": "ev-01: programme launches",
             "scope": None, "shifted_dimensions": ["object"], "reasoning": "Structured onboarding transfers."},
        ],
        "unknowns": [],
    })


def run() -> None:
    pre_existing = {
        p: Path(p).exists()
        for p in ("data/raw", "data/jobs", "data/evaluations", "data/runs")
    }
    for path, existed in pre_existing.items():
        check(not existed, f"expected {path} to not exist yet -- refusing to run a destructive cleanup test otherwise")
    if any(pre_existing.values()):
        print("Aborting: one or more data/ paths already exist. Not safe to auto-clean; inspect manually.")
        raise SystemExit(1)

    orchestrator_module.SOURCE_ADAPTERS["Indeed"] = FakeAdapter

    candidate_profile = CandidateProfile.load()
    manifest = RunManifest.new(
        sourcing_config_version="sourcing-config-v0.1",
        candidate_profile_version=candidate_profile.version,
        hard_disqualifier_policy_version="hard-disqualifier-v0.2",
        capability_fit_policy_version="capability-fit-v0.3",
        extraction_prompt_version="job-extraction-v0.1",
        output_schema_version="eval-result-v0.2",
        model_config_version="mvp-model-config-v0.1",
        sources_enabled=["Indeed"],
        max_jobs_per_source=2,
    )

    orchestrator = PipelineOrchestrator(
        manifest=manifest, candidate_profile=candidate_profile,
        query_catalog=QueryCatalog(), location=LocationFilter(city="Munich", country_code="DE"),
        interactive=False,
    )
    # Swap the wrapped raw SDK client on each stage's ModelClient -- the
    # ModelClient interface itself (system_blocks/user_message in, text/usage
    # out) is exercised for real; only the underlying Anthropic call is faked.
    orchestrator.extractor.model_client._client = FakeClient(_fake_extraction_response)
    orchestrator.hd_evaluator.model_client._client = FakeClient(_fake_hd_response)
    orchestrator.fit_evaluator.model_client._client = FakeClient(_fake_fit_response)

    try:
        orchestrator.run()

        c = orchestrator.reconciler.counts
        check(c.discovered == 2, f"expected discovered=2, got {c.discovered}")
        check(c.captured == 2, f"expected captured=2, got {c.captured}")
        check(c.new == 2, f"expected new=2, got {c.new}")
        check(c.hard_disqualifier_excluded == 1, f"expected hard_disqualifier_excluded=1, got {c.hard_disqualifier_excluded}")
        check(c.fit_completed == 1, f"expected fit_completed=1, got {c.fit_completed}")
        check(c.needs_human_review == 1, f"expected needs_human_review=1 (Unknown HD rule on the pass job), got {c.needs_human_review}")
        check(not orchestrator.reconciler.reconcile(), f"expected no reconciliation problems, got {orchestrator.reconciler.reconcile()}")

        results = orchestrator.eval_repo.read_all_for_run(manifest.run_id)
        check(len(results) == 2,
              f"expected 2 evaluation results tagged with manifest.run_id={manifest.run_id}, got {len(results)} "
              "-- this is exactly the run_id-desync bug class this test exists to catch")

        fit_labels = [r.capability_fit.label for r in results if r.capability_fit.label]
        check(fit_labels == ["PlausibleFit"],
              f"expected one derived PlausibleFit label (core done + adjacent), got {fit_labels}")

        jobs = orchestrator.job_repo.read_current()
        check(len(jobs) == 2, f"expected 2 persisted Job records, got {len(jobs)}")

        export_path = Path(f"data/runs/{manifest.run_id}/review_export.xlsx")
        check(export_path.exists(), f"expected review export at {export_path}")
        if export_path.exists():
            skill_map = load_workbook(export_path, rich_text=True)["Fit Skill Map"]
            rows = [[c.value for c in r] for r in skill_map.iter_rows(min_row=2)]
            requirements = [r[5] for r in rows]
            check(requirements == ["★ B2B account management (required)",
                                   "★ Customer onboarding (responsibility)"],
                  f"skill map should list the Fit job's skills, core first, got {requirements}")
            if len(rows) == 2:
                check(str(rows[0][7]) == "Direct evidence. Owned B2B accounts end-to-end.",
                      f"done row should render the fixed verdict + reasoning, got {rows[0][7]!r}")
                check(str(rows[1][7]).startswith("Transferable (shifted: object)."),
                      f"adjacent row should name its shifted dimensions, got {rows[1][7]!r}")
                check("Scope: Owned two named accounts." in rows[0][6],
                      f"done row evidence should carry the scope, got {rows[0][6]!r}")

    finally:
        for path in ("data/raw", "data/jobs", "data/evaluations", "data/runs"):
            shutil.rmtree(path, ignore_errors=True)

    if problems:
        print(f"{len(problems)} PROBLEMS:")
        for p in problems:
            print(f"- {p}")
        raise SystemExit(1)
    print("Phase D smoke test: all checks passed, 0 problems, 0 live API/Apify calls made.")


if __name__ == "__main__":
    run()
